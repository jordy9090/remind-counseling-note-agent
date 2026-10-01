"""Formatted text from the final document editor: parsing and DOCX/PDF export. Synthetic text only."""
from __future__ import annotations

import unittest
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from app.schemas.document import DocumentExportRequest
from app.services.document_export import (
    DocumentExportService,
    render_pdf_with_reportlab,
    render_sections_html,
    renderable_sections,
)
from app.services.rich_text import parse_color, parse_rich_text

MARKUP = (
    '<div><b>굵게</b> 보통 <i>기울임</i> <u>밑줄</u></div>'
    '<div style="text-align: center;"><font color="#dc2626">빨강</font>'
    '<span style="background-color: rgb(254, 240, 138);">형광</span></div>'
    '<div><br></div>'
    '<ul><li>첫째</li><li><strong>둘째</strong></li></ul>'
    '<ol><li>하나</li><li>둘</li></ol>'
)


def export_request(fmt: str, **section) -> DocumentExportRequest:
    return DocumentExportRequest.model_validate({
        "format": fmt, "document_type": "session_note", "case_id": "SYNTHETIC-RICH", "session_number": 1,
        "session_date": "2026-10-02", "title": "상담일지",
        "sections": [{"id": "s1", "title": "상담 내용", **section}],
    })


class RichTextParsingTest(unittest.TestCase):
    def test_marks_colors_alignment_and_lists(self) -> None:
        paragraphs = parse_rich_text(MARKUP)
        self.assertEqual([p.text for p in paragraphs], ["굵게 보통 기울임 밑줄", "빨강형광", "", "첫째", "둘째", "하나", "둘"])
        first = paragraphs[0].runs
        self.assertTrue(first[0].bold and not first[0].italic)
        self.assertEqual((first[1].text, first[1].bold), (" 보통 ", False))
        self.assertTrue(first[2].italic)
        self.assertTrue(first[4].underline)
        self.assertEqual(paragraphs[1].align, "center")
        self.assertEqual(paragraphs[1].runs[0].color, "DC2626")
        self.assertEqual(paragraphs[1].runs[1].highlight, "FEF08A")
        self.assertEqual([(p.list_kind, p.number) for p in paragraphs[3:]],
                         [("bullet", 1), ("bullet", 2), ("number", 1), ("number", 2)])
        self.assertTrue(paragraphs[4].runs[0].bold)

    def test_css_marks_and_color_forms(self) -> None:
        run = parse_rich_text(
            '<span style="font-weight: 700; font-style: italic; text-decoration-line: underline; color: #0a0">x</span>'
        )[0].runs[0]
        self.assertTrue(run.bold and run.italic and run.underline)
        self.assertEqual(run.color, "00AA00")
        self.assertEqual(parse_color("rgba(1, 2, 3, 0.5)"), "010203")
        for rejected in ("red", "rgba(0, 0, 0, 0)", "rgb(300, 0, 0)", "url(x)", "#12", "expression(1)", None):
            self.assertIsNone(parse_color(rejected))

    def test_unsafe_markup_is_reduced_to_text(self) -> None:
        paragraphs = parse_rich_text(
            '<div onclick="x()">안전<script>alert(1)</script><style>p{}</style>'
            '<img src=x onerror=alert(1)><a href="javascript:alert(1)">링크</a>'
            '<span style="color: red; background-image: url(http://x)">색</span> &lt;b&gt;</div>'
        )
        self.assertEqual([p.text for p in paragraphs], ["안전링크색 <b>"])
        self.assertTrue(all(run.color is None and run.highlight is None and not run.bold for run in paragraphs[0].runs))
        self.assertEqual(parse_rich_text(""), [])
        self.assertEqual(parse_rich_text("<div><br></div><div><br></div>"), [])


class RichTextExportTest(unittest.TestCase):
    def test_docx_keeps_the_formatting(self) -> None:
        result = DocumentExportService().export(export_request("docx", content="plain", content_html=MARKUP))
        document = Document(BytesIO(result.content))
        by_text = {p.text: p for p in document.paragraphs}
        runs = {run.text: run for run in by_text["굵게 보통 기울임 밑줄"].runs}
        self.assertTrue(runs["굵게"].bold)
        self.assertTrue(runs["기울임"].italic)
        self.assertTrue(runs["밑줄"].underline)
        self.assertFalse(runs[" 보통 "].bold)
        centered = by_text["빨강형광"]
        self.assertEqual(centered.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(str(centered.runs[0].font.color.rgb), "DC2626")
        shading = centered.runs[1]._element.rPr.find(qn("w:shd"))
        self.assertEqual(shading.get(qn("w:fill")), "FEF08A")
        self.assertEqual(by_text["첫째"].style.name, "List Bullet")
        self.assertIn("1. 하나", by_text)
        self.assertIn("2. 둘", by_text)
        self.assertNotIn("plain", by_text)

    def test_plain_sections_are_exported_as_before(self) -> None:
        for extra in ({}, {"content_html": ""}, {"content_html": "<div><br></div>"}):
            document = Document(BytesIO(DocumentExportService().export(
                export_request("docx", content="첫 줄\n• 항목", **extra)).content))
            texts = [p.text for p in document.paragraphs]
            self.assertIn("첫 줄", texts)
            self.assertEqual(next(p for p in document.paragraphs if p.text == "항목").style.name, "List Bullet")

    def test_list_sections_and_typed_prefixes_stay_lists(self) -> None:
        document = Document(BytesIO(DocumentExportService().export(
            export_request("docx", content=["가", "나"], content_html="<div><b>가</b></div><div>- 나</div><div>3) 다</div>")
        ).content))
        by_text = {p.text: p for p in document.paragraphs}
        self.assertEqual(by_text["가"].style.name, "List Bullet")
        self.assertTrue(by_text["가"].runs[0].bold)
        self.assertEqual(by_text["나"].style.name, "List Bullet")
        self.assertIn("3. 다", by_text)

    def test_supervision_blocks_use_formatted_text_except_tables_and_transcripts(self) -> None:
        request = DocumentExportRequest.model_validate({
            "format": "docx", "document_type": "supervision_report", "case_id": "SYNTHETIC-RICH", "session_number": 1,
            "title": "수퍼비전 보고서",
            "sections": [{"id": "s", "title": "성찰", "level": 2, "contentBlocks": [
                {"id": "b1", "type": "paragraph", "text": "성찰 내용", "textHtml": "<div><u>성찰</u> 내용</div>"},
                {"id": "b2", "type": "reflection_box", "text": "상자", "text_html": '<div style="text-align:right"><b>상자</b></div>'},
                {"id": "b3", "type": "table", "rows": [{"회기": "1"}], "text_html": "<b>무시</b>"},
                {"id": "b4", "type": "paragraph", "text": "[상담사 확인 필요]", "text_html": "<div>[상담사 확인 필요]</div>"},
            ]}],
        })
        document = Document(BytesIO(DocumentExportService().export(request).content))
        paragraph = next(p for p in document.paragraphs if p.text == "성찰 내용")
        self.assertTrue(paragraph.runs[0].underline)
        box = next(table.cell(0, 0) for table in document.tables if table.cell(0, 0).text == "상자")
        self.assertEqual(box.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertTrue(box.paragraphs[0].runs[0].bold)
        all_text = "\n".join(p.text for p in document.paragraphs)
        self.assertNotIn("무시", all_text)
        self.assertNotIn("확인 필요", all_text)

    def test_pdf_renderers_accept_formatted_text(self) -> None:
        request = export_request("pdf", content="plain", content_html=MARKUP + "<div>a &amp; b &lt;script&gt;</div>")
        pdf = render_pdf_with_reportlab(request)
        self.assertTrue(pdf.startswith(b"%PDF"))
        # Bold switches to the Gothic face because the built-in Korean fonts have one weight.
        self.assertIn(b"HYGothic-Medium", pdf)
        html_body = render_sections_html(renderable_sections(request.sections))
        self.assertIn("<strong>굵게</strong>", html_body)
        self.assertIn('<p style="text-align:center"><span style="color:#DC2626">빨강</span>', html_body)
        self.assertIn('<span style="background-color:#FEF08A">형광</span>', html_body)
        self.assertIn("<ul><li>첫째</li></ul>", html_body)
        self.assertIn("<p>1. 하나</p>", html_body)
        self.assertIn("a &amp; b &lt;script&gt;", html_body)
        self.assertNotIn("<script", html_body)

    def test_markup_length_is_bounded(self) -> None:
        with self.assertRaises(ValueError):
            export_request("docx", content="x", content_html="<b>" + "가" * 200_001 + "</b>")


if __name__ == "__main__":
    unittest.main()
